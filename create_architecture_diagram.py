"""
create_architecture_diagram.py - Generate a visual architecture diagram Word document.

Uses python-docx tables with colored cells to create box-diagram style
architecture visuals. No external image dependencies.

Usage:
    python create_architecture_diagram.py

Output:
    docs/ClaimClear_AI_Architecture_Diagram.docx
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ═══════════════════════════════════════════════════════════════════════════
# Helpers
# ═══════════════════════════════════════════════════════════════════════════

NAVY = "0F4C81"
TEAL = "17A2B8"
LIGHT_BLUE = "D6EAF8"
LIGHT_TEAL = "D1F2EB"
LIGHT_GREEN = "D5F5E3"
LIGHT_YELLOW = "FEF9E7"
LIGHT_ORANGE = "FDEBD0"
LIGHT_PURPLE = "E8DAEF"
LIGHT_GRAY = "F2F3F4"
WHITE = "FFFFFF"
DARK_TEXT = "2C3E50"


def set_cell_bg(cell, color: str):
    """Set cell background color."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text: str, bold=False, size=10, color=DARK_TEXT, align="center"):
    """Set cell text with formatting."""
    cell.text = ""
    para = cell.paragraphs[0]
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_spacer(doc, height=0.2):
    """Add vertical space."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Cm(height)
    para.paragraph_format.space_after = Cm(0)


def add_arrow_text(doc, text="▼", size=14):
    """Add a centered arrow/connector."""
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    run.font.bold = True
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(2)


def add_section_heading(doc, text: str, level=2):
    """Add a colored heading."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor.from_string(NAVY)


def add_subsection_heading(doc, text: str):
    """Add a teal subsection heading."""
    heading = doc.add_heading(text, level=3)
    for run in heading.runs:
        run.font.color.rgb = RGBColor.from_string(TEAL)


def add_body(doc, text: str):
    """Add body paragraph."""
    para = doc.add_paragraph(text)
    for run in para.runs:
        run.font.size = Pt(10)


def make_box_table(doc, rows_data, col_widths=None):
    """Create a table representing a box diagram component."""
    num_cols = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for ri, row in enumerate(rows_data):
        for ci, cell_data in enumerate(row):
            if ci >= num_cols:
                continue
            cell = table.cell(ri, ci)
            text = cell_data.get("text", "")
            bg = cell_data.get("bg", WHITE)
            bold = cell_data.get("bold", False)
            size = cell_data.get("size", 10)
            color = cell_data.get("color", DARK_TEXT)
            align = cell_data.get("align", "center")
            colspan = cell_data.get("colspan", 1)

            set_cell_bg(cell, bg)
            set_cell_text(cell, text, bold=bold, size=size, color=color, align=align)

            # Handle colspan by merging cells
            if colspan > 1 and ci + colspan - 1 < num_cols:
                for merge_ci in range(ci + 1, ci + colspan):
                    cell.merge(table.cell(ri, merge_ci))

    return table


# ═══════════════════════════════════════════════════════════════════════════
# Document Creation
# ═══════════════════════════════════════════════════════════════════════════

def create_diagram():
    doc = Document()

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor.from_string(DARK_TEXT)

    # ── Title ──
    title = doc.add_heading("ClaimClear AI", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor.from_string(NAVY)

    subtitle = doc.add_paragraph("Architecture Diagram")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor.from_string(TEAL)
        run.font.italic = True

    doc.add_paragraph("")

    # ══════════════════════════════════════════════════════════════════════
    # DIAGRAM 1: High-Level System Architecture
    # ══════════════════════════════════════════════════════════════════════

    add_section_heading(doc, "1. High-Level System Architecture")
    add_body(doc, "The system consists of four main layers: User Interface, Agentic Pipeline, RAG Knowledge Store, and LLM Provider.")

    # User layer
    make_box_table(doc, [[
        {"text": "User (Web Browser)", "bg": LIGHT_BLUE, "bold": True, "size": 12, "colspan": 3},
        {"text": "", "bg": LIGHT_BLUE}, {"text": "", "bg": LIGHT_BLUE},
    ]])
    add_arrow_text(doc, "| HTTP :8501 |")
    add_arrow_text(doc, "▼")

    # Streamlit layer
    make_box_table(doc, [
        [
            {"text": "Streamlit UI  (app.py)", "bg": NAVY, "bold": True, "size": 11, "color": WHITE, "colspan": 3},
            {"text": "", "bg": NAVY}, {"text": "", "bg": NAVY},
        ],
        [
            {"text": "Claim Input Form\n(2-column layout)", "bg": LIGHT_BLUE, "size": 9},
            {"text": "Sidebar\n(Model, API Key,\nTone, RAG settings)", "bg": LIGHT_BLUE, "size": 9},
            {"text": "Results Display\n(Explanation, Scores,\nRAG Citations, Glossary)", "bg": LIGHT_BLUE, "size": 9},
        ],
    ])
    add_arrow_text(doc, "▼")

    # Pipeline layer
    make_box_table(doc, [
        [
            {"text": "Agentic Pipeline  (pipeline.py)", "bg": NAVY, "bold": True, "size": 11, "color": WHITE, "colspan": 5},
            {"text": "", "bg": NAVY}, {"text": "", "bg": NAVY}, {"text": "", "bg": NAVY}, {"text": "", "bg": NAVY},
        ],
        [
            {"text": "1. ANALYZE\nExtract factors\n& complexity", "bg": LIGHT_GREEN, "size": 8, "bold": True},
            {"text": "2. RETRIEVE\nRAG from\nPolicy PDF", "bg": LIGHT_TEAL, "size": 8, "bold": True},
            {"text": "3. GENERATE\nExplanation\n+ Glossary", "bg": LIGHT_YELLOW, "size": 8, "bold": True},
            {"text": "4. EVALUATE\nScore 4\nDimensions", "bg": LIGHT_ORANGE, "size": 8, "bold": True},
            {"text": "5. REFINE\n(if < 7/10)\nFix Issues", "bg": LIGHT_PURPLE, "size": 8, "bold": True},
        ],
    ])
    add_arrow_text(doc, "▼                        ▼")

    # Two boxes side by side: RAG and LLM
    make_box_table(doc, [
        [
            {"text": "PolicyStore  (policy_store.py)", "bg": TEAL, "bold": True, "size": 10, "color": WHITE, "colspan": 2},
            {"text": "", "bg": TEAL},
            {"text": "LLM Provider", "bg": NAVY, "bold": True, "size": 10, "color": WHITE, "colspan": 2},
            {"text": "", "bg": NAVY},
        ],
        [
            {"text": "LangChain\nPyPDFLoader\nTextSplitter", "bg": LIGHT_TEAL, "size": 8},
            {"text": "ChromaDB\nVector Store\n+ Keyword FB", "bg": LIGHT_TEAL, "size": 8},
            {"text": "", "bg": WHITE, "size": 8},
            {"text": "OpenAI\nGPT-4o-mini\nGPT-4o", "bg": LIGHT_BLUE, "size": 8},
            {"text": "TCS GenAI\nLab / Custom\nEndpoint", "bg": LIGHT_BLUE, "size": 8},
        ],
    ])
    add_arrow_text(doc, "▲")

    # PDF source
    make_box_table(doc, [[
        {"text": "SilverShield_Master_Policy.pdf  (8 pages, 12 sections)", "bg": LIGHT_GRAY, "bold": True, "size": 10, "colspan": 3},
        {"text": "", "bg": LIGHT_GRAY}, {"text": "", "bg": LIGHT_GRAY},
    ]])

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════
    # DIAGRAM 2: RAG Pipeline Detail
    # ══════════════════════════════════════════════════════════════════════

    add_section_heading(doc, "2. RAG Pipeline Architecture")
    add_body(doc, "The RAG (Retrieval-Augmented Generation) pipeline ingests the policy PDF, splits it into chunks, and retrieves relevant sections with page-level citations.")

    # PDF source
    make_box_table(doc, [[
        {"text": "SilverShield_Master_Policy.pdf", "bg": LIGHT_GRAY, "bold": True, "size": 11, "colspan": 2},
        {"text": "", "bg": LIGHT_GRAY},
    ]])
    add_arrow_text(doc, "▼")

    # PDF Loader
    make_box_table(doc, [[
        {"text": "LangChain PyPDFLoader", "bg": LIGHT_TEAL, "bold": True, "size": 10, "colspan": 2},
        {"text": "", "bg": LIGHT_TEAL},
    ], [
        {"text": "Extracts text from each page\nPreserves page numbers as metadata", "bg": LIGHT_TEAL, "size": 9, "colspan": 2},
        {"text": "", "bg": LIGHT_TEAL},
    ]])
    add_arrow_text(doc, "▼  (8 pages of raw text)")

    # Splitter
    make_box_table(doc, [[
        {"text": "RecursiveCharacterTextSplitter", "bg": LIGHT_GREEN, "bold": True, "size": 10, "colspan": 2},
        {"text": "", "bg": LIGHT_GREEN},
    ], [
        {"text": "chunk_size=500, chunk_overlap=80\nSeparators: paragraph > newline > sentence > word", "bg": LIGHT_GREEN, "size": 9, "colspan": 2},
        {"text": "", "bg": LIGHT_GREEN},
    ]])
    add_arrow_text(doc, "▼  (~18-25 chunks with page metadata)")

    # Dual path
    make_box_table(doc, [
        [
            {"text": "WITH API Key", "bg": NAVY, "bold": True, "size": 10, "color": WHITE},
            {"text": "", "bg": WHITE, "size": 8},
            {"text": "WITHOUT API Key", "bg": TEAL, "bold": True, "size": 10, "color": WHITE},
        ],
        [
            {"text": "OpenAI Embeddings\ntext-embedding-3-small\n      |\nChromaDB Vector Store\nsimilarity_search(k=3)", "bg": LIGHT_BLUE, "size": 9},
            {"text": "OR", "bg": WHITE, "bold": True, "size": 12, "color": TEAL},
            {"text": "Keyword Extraction\nTF-IDF Token Matching\n      |\nPolicy-Type Boost (2x)\nTop-3 by Score", "bg": LIGHT_YELLOW, "size": 9},
        ],
    ])
    add_arrow_text(doc, "▼")

    # RAG Context output
    make_box_table(doc, [[
        {"text": "RAGContext Output", "bg": NAVY, "bold": True, "size": 10, "color": WHITE, "colspan": 3},
        {"text": "", "bg": NAVY}, {"text": "", "bg": NAVY},
    ], [
        {"text": "Chunk 1\n[PDF, p.3]\nSection 3.2\nScore: 8.0", "bg": LIGHT_BLUE, "size": 8},
        {"text": "Chunk 2\n[PDF, p.4]\nSection 5.1\nScore: 6.5", "bg": LIGHT_BLUE, "size": 8},
        {"text": "Chunk 3\n[PDF, p.3]\nSection 3.3\nScore: 5.0", "bg": LIGHT_BLUE, "size": 8},
    ]])
    add_arrow_text(doc, "▼  (Injected into LLM prompt as cited context)")

    # LLM prompt
    make_box_table(doc, [[
        {"text": "LLM Generate Stage - Prompt includes:", "bg": LIGHT_YELLOW, "bold": True, "size": 10, "colspan": 2},
        {"text": "", "bg": LIGHT_YELLOW},
    ], [
        {"text": "[SilverShield_Master_Policy.pdf, p.3, Section 3.2]: Out-of-network\nservices are covered at a reduced rate under PPO plans...", "bg": LIGHT_YELLOW, "size": 8, "colspan": 2, "align": "left"},
        {"text": "", "bg": LIGHT_YELLOW},
    ]])

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════
    # DIAGRAM 3: Pipeline Stage Detail
    # ══════════════════════════════════════════════════════════════════════

    add_section_heading(doc, "3. Agentic Pipeline Stage Flow")
    add_body(doc, "Each stage is a focused LLM call with structured JSON output. The pipeline is sequential with conditional branching at the evaluation gate.")

    # Stage 1
    make_box_table(doc, [
        [{"text": "STAGE 1: ANALYZE", "bg": LIGHT_GREEN, "bold": True, "size": 11, "colspan": 2}, {"text": "", "bg": LIGHT_GREEN}],
        [{"text": "Input: Claim data (ID, type, amount, decision, reason)", "bg": LIGHT_GREEN, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_GREEN}],
        [{"text": "LLM Task: Extract complexity, key factors, jargon terms", "bg": LIGHT_GREEN, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_GREEN}],
        [{"text": "Output: AnalysisResult (JSON)", "bg": LIGHT_GREEN, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_GREEN}],
        [{"text": "Tokens: ~200  |  Temperature: 0.3", "bg": LIGHT_GREEN, "size": 8, "colspan": 2}, {"text": "", "bg": LIGHT_GREEN}],
    ])
    add_arrow_text(doc, "▼")

    # Stage 2
    make_box_table(doc, [
        [{"text": "STAGE 2: RETRIEVE (RAG)", "bg": LIGHT_TEAL, "bold": True, "size": 11, "colspan": 2}, {"text": "", "bg": LIGHT_TEAL}],
        [{"text": "Input: Claim query (type + decision + reason)", "bg": LIGHT_TEAL, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_TEAL}],
        [{"text": "Action: Search PolicyStore (ChromaDB or keyword)", "bg": LIGHT_TEAL, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_TEAL}],
        [{"text": "Source: SilverShield_Master_Policy.pdf (8 pages)", "bg": LIGHT_TEAL, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_TEAL}],
        [{"text": "Output: RAGContext (top-3 chunks with page citations)", "bg": LIGHT_TEAL, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_TEAL}],
        [{"text": "No LLM call  |  Vector or Keyword search", "bg": LIGHT_TEAL, "size": 8, "colspan": 2}, {"text": "", "bg": LIGHT_TEAL}],
    ])
    add_arrow_text(doc, "▼")

    # Stage 3
    make_box_table(doc, [
        [{"text": "STAGE 3: GENERATE", "bg": LIGHT_YELLOW, "bold": True, "size": 11, "colspan": 2}, {"text": "", "bg": LIGHT_YELLOW}],
        [{"text": "Input: Claim + Analysis + RAG Context + Few-Shot Example", "bg": LIGHT_YELLOW, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_YELLOW}],
        [{"text": "LLM Task: Write personalized explanation letter with citations", "bg": LIGHT_YELLOW, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_YELLOW}],
        [{"text": "Output: Explanation + Glossary (JSON)", "bg": LIGHT_YELLOW, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_YELLOW}],
        [{"text": "Tokens: ~600  |  Includes PDF page citations", "bg": LIGHT_YELLOW, "size": 8, "colspan": 2}, {"text": "", "bg": LIGHT_YELLOW}],
    ])
    add_arrow_text(doc, "▼")

    # Stage 4
    make_box_table(doc, [
        [{"text": "STAGE 4: EVALUATE", "bg": LIGHT_ORANGE, "bold": True, "size": 11, "colspan": 2}, {"text": "", "bg": LIGHT_ORANGE}],
        [{"text": "Input: Original claim + Generated explanation", "bg": LIGHT_ORANGE, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_ORANGE}],
        [{"text": "LLM Task: Score accuracy, empathy, readability, completeness (1-10)", "bg": LIGHT_ORANGE, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_ORANGE}],
        [{"text": "Output: EvaluationResult + issues + suggestions (JSON)", "bg": LIGHT_ORANGE, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_ORANGE}],
        [{"text": "Tokens: ~300  |  Quality gate: score >= 7 passes", "bg": LIGHT_ORANGE, "size": 8, "colspan": 2}, {"text": "", "bg": LIGHT_ORANGE}],
    ])
    add_arrow_text(doc, "▼  Score < 7?")

    # Decision gate
    make_box_table(doc, [
        [
            {"text": "YES (Score < 7)", "bg": LIGHT_PURPLE, "bold": True, "size": 10},
            {"text": "NO (Score >= 7)", "bg": LIGHT_GREEN, "bold": True, "size": 10},
        ],
        [
            {"text": "Proceed to Stage 5\nRefine explanation", "bg": LIGHT_PURPLE, "size": 9},
            {"text": "Skip refinement\nReturn result directly", "bg": LIGHT_GREEN, "size": 9},
        ],
    ])
    add_arrow_text(doc, "▼ (conditional)")

    # Stage 5
    make_box_table(doc, [
        [{"text": "STAGE 5: REFINE (Conditional)", "bg": LIGHT_PURPLE, "bold": True, "size": 11, "colspan": 2}, {"text": "", "bg": LIGHT_PURPLE}],
        [{"text": "Input: Explanation + Evaluation issues + Suggestions", "bg": LIGHT_PURPLE, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_PURPLE}],
        [{"text": "LLM Task: Fix specific issues identified in evaluation", "bg": LIGHT_PURPLE, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_PURPLE}],
        [{"text": "Output: Improved explanation + glossary (JSON)", "bg": LIGHT_PURPLE, "size": 9, "align": "left", "colspan": 2}, {"text": "", "bg": LIGHT_PURPLE}],
        [{"text": "Tokens: ~500  |  Only runs when needed", "bg": LIGHT_PURPLE, "size": 8, "colspan": 2}, {"text": "", "bg": LIGHT_PURPLE}],
    ])
    add_arrow_text(doc, "▼")

    # Final output
    make_box_table(doc, [[
        {"text": "PipelineResult", "bg": NAVY, "bold": True, "size": 11, "color": WHITE, "colspan": 3},
        {"text": "", "bg": NAVY}, {"text": "", "bg": NAVY},
    ], [
        {"text": "Explanation\n+ Glossary\n+ PDF Citations", "bg": LIGHT_BLUE, "size": 9},
        {"text": "Quality Scores\n(4 dimensions)\n+ Issues", "bg": LIGHT_BLUE, "size": 9},
        {"text": "RAG Context\n(chunks, pages,\nmethod used)", "bg": LIGHT_BLUE, "size": 9},
    ]])

    doc.add_page_break()

    # ══════════════════════════════════════════════════════════════════════
    # DIAGRAM 4: Technology Stack
    # ══════════════════════════════════════════════════════════════════════

    add_section_heading(doc, "4. Technology Stack")

    make_box_table(doc, [
        [
            {"text": "Layer", "bg": NAVY, "bold": True, "size": 10, "color": WHITE},
            {"text": "Technology", "bg": NAVY, "bold": True, "size": 10, "color": WHITE},
            {"text": "Purpose", "bg": NAVY, "bold": True, "size": 10, "color": WHITE},
        ],
        [
            {"text": "Frontend", "bg": LIGHT_BLUE, "bold": True, "size": 9},
            {"text": "Streamlit 1.41+", "bg": LIGHT_BLUE, "size": 9},
            {"text": "Web UI with forms,\ncharts, sidebar", "bg": LIGHT_BLUE, "size": 9},
        ],
        [
            {"text": "Pipeline", "bg": LIGHT_GREEN, "bold": True, "size": 9},
            {"text": "Python + OpenAI SDK", "bg": LIGHT_GREEN, "size": 9},
            {"text": "5-stage agentic\norchestration", "bg": LIGHT_GREEN, "size": 9},
        ],
        [
            {"text": "RAG Framework", "bg": LIGHT_TEAL, "bold": True, "size": 9},
            {"text": "LangChain", "bg": LIGHT_TEAL, "size": 9},
            {"text": "PDF loading, text\nsplitting, embeddings", "bg": LIGHT_TEAL, "size": 9},
        ],
        [
            {"text": "Vector Store", "bg": LIGHT_TEAL, "bold": True, "size": 9},
            {"text": "ChromaDB (in-memory)", "bg": LIGHT_TEAL, "size": 9},
            {"text": "Semantic search with\nOpenAI embeddings", "bg": LIGHT_TEAL, "size": 9},
        ],
        [
            {"text": "Document", "bg": LIGHT_GRAY, "bold": True, "size": 9},
            {"text": "SilverShield PDF\n(8 pages, 12 sections)", "bg": LIGHT_GRAY, "size": 9},
            {"text": "RAG source document\nwith page citations", "bg": LIGHT_GRAY, "size": 9},
        ],
        [
            {"text": "LLM", "bg": LIGHT_YELLOW, "bold": True, "size": 9},
            {"text": "OpenAI GPT-4o-mini\nGPT-4o / TCS GenAI", "bg": LIGHT_YELLOW, "size": 9},
            {"text": "Structured JSON\ngeneration + eval", "bg": LIGHT_YELLOW, "size": 9},
        ],
        [
            {"text": "Embeddings", "bg": LIGHT_PURPLE, "bold": True, "size": 9},
            {"text": "text-embedding-3-small", "bg": LIGHT_PURPLE, "size": 9},
            {"text": "Vector embeddings\nfor semantic RAG", "bg": LIGHT_PURPLE, "size": 9},
        ],
    ])

    add_spacer(doc, 1.0)

    # ══════════════════════════════════════════════════════════════════════
    # DIAGRAM 5: Data Flow
    # ══════════════════════════════════════════════════════════════════════

    add_section_heading(doc, "5. Data Flow Summary")

    make_box_table(doc, [
        [
            {"text": "Step", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
            {"text": "Action", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
            {"text": "Data", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
            {"text": "Tokens", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
        ],
        [
            {"text": "1", "bg": LIGHT_BLUE, "size": 9, "bold": True},
            {"text": "User fills claim form", "bg": LIGHT_BLUE, "size": 9},
            {"text": "ClaimInput dataclass", "bg": LIGHT_BLUE, "size": 9},
            {"text": "-", "bg": LIGHT_BLUE, "size": 9},
        ],
        [
            {"text": "2", "bg": LIGHT_GREEN, "size": 9, "bold": True},
            {"text": "Analyze claim", "bg": LIGHT_GREEN, "size": 9},
            {"text": "AnalysisResult (JSON)", "bg": LIGHT_GREEN, "size": 9},
            {"text": "~200", "bg": LIGHT_GREEN, "size": 9},
        ],
        [
            {"text": "3", "bg": LIGHT_TEAL, "size": 9, "bold": True},
            {"text": "RAG retrieval from PDF", "bg": LIGHT_TEAL, "size": 9},
            {"text": "RAGContext (3 chunks)", "bg": LIGHT_TEAL, "size": 9},
            {"text": "0*", "bg": LIGHT_TEAL, "size": 9},
        ],
        [
            {"text": "4", "bg": LIGHT_YELLOW, "size": 9, "bold": True},
            {"text": "Generate explanation", "bg": LIGHT_YELLOW, "size": 9},
            {"text": "Explanation + Glossary", "bg": LIGHT_YELLOW, "size": 9},
            {"text": "~600", "bg": LIGHT_YELLOW, "size": 9},
        ],
        [
            {"text": "5", "bg": LIGHT_ORANGE, "size": 9, "bold": True},
            {"text": "Evaluate quality", "bg": LIGHT_ORANGE, "size": 9},
            {"text": "EvaluationResult (JSON)", "bg": LIGHT_ORANGE, "size": 9},
            {"text": "~300", "bg": LIGHT_ORANGE, "size": 9},
        ],
        [
            {"text": "6", "bg": LIGHT_PURPLE, "size": 9, "bold": True},
            {"text": "Refine (if score < 7)", "bg": LIGHT_PURPLE, "size": 9},
            {"text": "Improved explanation", "bg": LIGHT_PURPLE, "size": 9},
            {"text": "~500", "bg": LIGHT_PURPLE, "size": 9},
        ],
        [
            {"text": "7", "bg": LIGHT_BLUE, "size": 9, "bold": True},
            {"text": "Display results + citations", "bg": LIGHT_BLUE, "size": 9},
            {"text": "PipelineResult", "bg": LIGHT_BLUE, "size": 9},
            {"text": "-", "bg": LIGHT_BLUE, "size": 9},
        ],
    ])

    add_body(doc, "* RAG retrieval uses no LLM tokens (keyword mode) or embedding tokens only (vector mode).")
    add_body(doc, "Typical total: ~1,100 tokens (no refinement) or ~1,600 tokens (with refinement).")

    add_spacer(doc, 1.0)

    # ══════════════════════════════════════════════════════════════════════
    # DIAGRAM 6: Multi-Model Support
    # ══════════════════════════════════════════════════════════════════════

    add_section_heading(doc, "6. Multi-Model Provider Architecture")

    make_box_table(doc, [
        [
            {"text": "Provider", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
            {"text": "Model ID", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
            {"text": "Base URL", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
            {"text": "API Key", "bg": NAVY, "bold": True, "size": 9, "color": WHITE},
        ],
        [
            {"text": "OpenAI", "bg": LIGHT_BLUE, "size": 9, "bold": True},
            {"text": "gpt-4o-mini", "bg": LIGHT_BLUE, "size": 9},
            {"text": "api.openai.com (default)", "bg": LIGHT_BLUE, "size": 9},
            {"text": "Required", "bg": LIGHT_BLUE, "size": 9},
        ],
        [
            {"text": "OpenAI", "bg": LIGHT_BLUE, "size": 9, "bold": True},
            {"text": "gpt-4o", "bg": LIGHT_BLUE, "size": 9},
            {"text": "api.openai.com (default)", "bg": LIGHT_BLUE, "size": 9},
            {"text": "Required", "bg": LIGHT_BLUE, "size": 9},
        ],
        [
            {"text": "TCS GenAI Lab", "bg": LIGHT_GREEN, "size": 9, "bold": True},
            {"text": "azure/genailab-\nmaas-gpt-4o", "bg": LIGHT_GREEN, "size": 9},
            {"text": "genailab.tcs.in", "bg": LIGHT_GREEN, "size": 9},
            {"text": "Not needed", "bg": LIGHT_GREEN, "size": 9},
        ],
        [
            {"text": "Custom", "bg": LIGHT_YELLOW, "size": 9, "bold": True},
            {"text": "Any", "bg": LIGHT_YELLOW, "size": 9},
            {"text": "User-provided URL", "bg": LIGHT_YELLOW, "size": 9},
            {"text": "Depends", "bg": LIGHT_YELLOW, "size": 9},
        ],
    ])

    add_spacer(doc, 0.5)
    add_body(doc, "All providers use the OpenAI-compatible API format via the OpenAI Python SDK. The get_openai_client() factory configures the correct base_url and API key for each provider.")

    # ── Footer ──
    doc.add_paragraph("")
    footer = doc.add_paragraph("ClaimClear AI - Architecture Diagram Document")
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("999999")
        run.font.italic = True

    # Save
    os.makedirs("docs", exist_ok=True)
    output_path = "docs/ClaimClear_AI_Architecture_Diagram.docx"
    doc.save(output_path)
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    create_diagram()
